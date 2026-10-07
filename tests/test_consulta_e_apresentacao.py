"""Consulta por descrição, apresentação fiscal e validação da planilha."""

from datetime import date

import openpyxl
import pandas as pd
import pytest

from src.busca import consultar
from src.carregamento import CAMINHO_BASE, COLUNAS_OBRIGATORIAS, ErroBase, carregar_base
from src.formatacao import EXIBICAO_BENEFICIO, campos_preenchidos, formatar_percentual, vigencia_exibicao
from tests.test_consulta import HOJE, base, beneficio, consulta, regra


def _planilha_teste(caminho, *, ncm_ativo=True, aliquotas=None, beneficios=None):
    tabelas = {nome: pd.DataFrame(columns=sorted(colunas)) for nome, colunas in COLUNAS_OBRIGATORIAS.items()}
    tabelas["NCM"] = pd.DataFrame([{"ncm": "12345678", "descricao_oficial": "Produto sintético", "ativo": "SIM"}])
    if not ncm_ativo:
        tabelas["NCM"] = tabelas["NCM"].drop(columns="ativo")
    if aliquotas is not None:
        tabelas["Aliquotas"] = pd.DataFrame(aliquotas)
    if beneficios is not None:
        tabelas["Beneficios"] = pd.DataFrame(beneficios)
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        for nome, tabela in tabelas.items():
            tabela.to_excel(writer, sheet_name=nome, index=False)


def test_ativo_obrigatorio_na_aba_ncm(tmp_path):
    caminho = tmp_path / "sem_ativo.xlsx"
    _planilha_teste(caminho, ncm_ativo=False)
    with pytest.raises(ErroBase, match="NCM: ativo"):
        carregar_base(caminho)


def test_ncm_inativo_ignorado_e_ativo_normalizado():
    b = base()
    b.ncm.loc[0, "ativo"] = "  nAo  "
    assert consulta(b).descricao_oficial == ""
    b.ncm.loc[0, "ativo"] = "  sIm  "
    assert consulta(b).descricao_oficial == "Descrição de teste"


def test_exige_descricao_sim_sem_descricao_mostra_possibilidade():
    r = consulta(base([regra(exige_descricao=" SIM ")]))
    assert r.regra_priorizada is None
    assert [c.identificador for c in r.possibilidades_aliquota] == ["R1"]
    assert r.nivel_confianca == "MEDIA"
    assert "Informe a descrição" in " ".join(r.motivos_confianca)


def test_exige_descricao_sim_compativel_prioriza():
    b = base([regra(exige_descricao="SIM", descricao_regra="Coca Cola", palavras_chave="coca cola;guarana")])
    r = consulta(b, descricao="Coca Cola PET 2L")
    assert r.regra_priorizada.identificador == "R1"
    assert r.regra_priorizada.score_descricao >= 70
    assert r.nivel_confianca == "ALTA"


def test_exige_descricao_sim_pouco_compativel_exige_revisao():
    b = base([regra(exige_descricao="SIM", descricao_regra="Coca Cola", palavras_chave="guarana")])
    r = consulta(b, descricao="Parafuso de teste")
    assert r.regra_priorizada is None
    assert len(r.possibilidades_aliquota) == 1
    assert r.nivel_confianca == "REVISAO_NECESSARIA"


def test_exige_descricao_nao_sem_descricao_mantem_alta_em_exato():
    r = consulta(base([regra(exige_descricao=" naO ")]))
    assert r.regra_priorizada.identificador == "R1"
    assert r.nivel_confianca == "ALTA"


@pytest.mark.parametrize(("valor", "esperado"), [
    (19, "19,00%"), (7, "7,00%"), (9.8, "9,80%"),
    (0.1, "0,10%"), (0.05, "0,05%"), ("9,8", "9,80%"),
])
def test_percentual_em_pontos_percentuais(valor, esperado):
    assert formatar_percentual(valor) == esperado


def test_palavras_chave_separadas_e_comparadas_individualmente():
    b = base([regra(descricao_regra="Outro produto", palavras_chave="refrigerante;coca cola;guarana;soda")])
    c = consulta(b, descricao="Coca Cola PET 2L").candidatos_aliquota[0]
    assert c.palavras_chave_utilizadas == ["refrigerante", "coca cola", "guarana", "soda"]
    assert c.melhor_palavra_chave == "coca cola"
    assert c.score_melhor_palavra_chave > c.score_descricao_regra
    assert c.score_descricao == c.score_melhor_palavra_chave


def test_maior_score_pode_vir_da_descricao_regra():
    b = base([regra(descricao_regra="Coca Cola PET 2L", palavras_chave="parafuso;borracha")])
    c = consulta(b, descricao="Coca Cola PET 2L").candidatos_aliquota[0]
    assert c.score_descricao_regra > c.score_melhor_palavra_chave
    assert c.score_descricao == c.score_descricao_regra


@pytest.mark.parametrize(("aplicacao", "rotulo"), [
    ("UNICO", "Benefício único"),
    ("ALTERNATIVO", "Benefícios alternativos"),
    ("CUMULATIVO", "Benefícios cumulativos"),
])
def test_aplicacao_beneficio_preservada_sem_conclusao_fiscal(aplicacao, rotulo):
    b = base(beneficios=[beneficio(aplicacao=f" {aplicacao.lower()} ", grupo_beneficio="GRUPO_TESTE", exige_descricao="NAO")])
    r = consulta(b)
    c = r.beneficios_priorizados[0]
    assert c.aplicacao == aplicacao
    assert c.grupo_beneficio == "GRUPO_TESTE"
    assert not r.beneficios_pendentes
    campos = dict(campos_preenchidos(c, EXIBICAO_BENEFICIO))
    assert campos["Grupo do benefício"] == "GRUPO_TESTE"
    assert campos["Aplicação"] == rotulo


@pytest.mark.parametrize(("codigo", "rotulo"), [
    ("REDUCAO_BASE_CALCULO", "Redução da base de cálculo"),
    ("ISENCAO", "Isenção"),
    ("DIFERIMENTO", "Diferimento"),
    ("CREDITO_OUTORGADO", "Crédito outorgado"),
])
def test_tipo_beneficio_traduzido_sem_alterar_codigo(codigo, rotulo):
    candidato = consulta(base(beneficios=[beneficio(tipo_beneficio=codigo)])).beneficios_priorizados[0]
    assert candidato.dados["tipo_beneficio"] == codigo
    assert dict(campos_preenchidos(candidato, EXIBICAO_BENEFICIO))["Tipo de benefício"] == rotulo


def test_apresentacao_publica_do_beneficio_omite_percentuais():
    candidato = consulta(base(beneficios=[beneficio(
        tipo_beneficio="REDUCAO_BASE_CALCULO",
        cbenef="GO123456",
        escopo_operacao="INTERNA",
        condicoes="Exigência legal sintética a verificar",
        percentual_reducao_bc="70",
        carga_efetiva="7",
        credito_outorgado_percentual="5",
    )])).beneficios_priorizados[0]

    campos = dict(campos_preenchidos(candidato, EXIBICAO_BENEFICIO))
    assert campos["Tipo de benefício"] == "Redução da base de cálculo"
    assert campos["Condições"] == "Exigência legal sintética a verificar"
    assert campos["cBenef"] == "GO123456"
    assert campos["Operação prevista"] == "Operação interna"
    assert "Redução da Base de Cálculo" not in campos
    assert "Carga tributária efetiva" not in campos
    assert "Crédito outorgado" not in campos
    assert candidato.dados["percentual_reducao_bc"] == "70"
    assert candidato.dados["carga_efetiva"] == "7"
    assert candidato.dados["credito_outorgado_percentual"] == "5"


def test_beneficio_que_exige_descricao_permanece_possivel():
    r = consulta(base(beneficios=[beneficio(exige_descricao="SIM", aplicacao="UNICO", escopo_operacao="INTERNA")]))
    assert not r.beneficios_priorizados
    assert [c.identificador for c in r.beneficios_possiveis] == ["B1"]
    assert r.beneficios_pendentes
    assert r.nivel_confianca == "MEDIA"


def test_correspondencia_alta_nao_comprova_condicoes_do_beneficio():
    r = consulta(base(beneficios=[beneficio(aplicacao="UNICO", escopo_operacao="INTERNA",
                                         condicoes="Exigência legal sintética a verificar")]))
    assert r.nivel_confianca == "ALTA"
    assert "regra EXATA ativa na base" in r.motivos_confianca[0]
    assert r.beneficios_priorizados[0].dados["condicoes"] == "Exigência legal sintética a verificar"
    assert r.beneficios_priorizados[0].aplicabilidade == "CONDICIONAL"


@pytest.mark.parametrize(("inicio", "fim", "qualificacao"), [
    (None, None, "ativa na base"),
    (date(2020, 1, 1), None, "ativa na base"),
    (None, date(2030, 12, 31), "ativa na base"),
    (date(2020, 1, 1), date(2030, 12, 31), "vigente"),
])
def test_motivo_precisao_so_afirma_vigencia_com_intervalo_cadastrado(inicio, fim, qualificacao):
    resultado = consulta(base([regra(vigencia_inicio=inicio, vigencia_fim=fim)]))
    assert resultado.regra_priorizada.identificador == "R1"
    assert resultado.nivel_confianca == "ALTA"
    assert f"regra EXATA {qualificacao}" in resultado.motivos_confianca[0]


@pytest.mark.parametrize(("inicio", "fim", "esperado"), [
    (None, None, "Vigência específica não cadastrada."),
    (date(2020, 1, 1), None, "Início da vigência: 01/01/2020"),
    (None, date(2030, 12, 31), "Fim da vigência: 31/12/2030"),
    (date(2020, 1, 1), date(2030, 12, 31), "01/01/2020 até 31/12/2030"),
])
def test_vigencia_exibida_sem_inferir_limite(inicio, fim, esperado):
    candidato = consulta(base([regra(vigencia_inicio=inicio, vigencia_fim=fim)])).regra_priorizada
    assert vigencia_exibicao(candidato) == esperado


def test_loader_aceita_exige_descricao_com_espacos_e_caixa_variavel(tmp_path):
    caminho = tmp_path / "normalizada.xlsx"
    _planilha_teste(caminho, aliquotas=[regra(exige_descricao=" sIm ")])
    candidato = consultar(carregar_base(caminho), "12345678", data_referencia=HOJE).candidatos_aliquota[0]
    assert candidato.exige_descricao == "SIM"


def test_duas_regras_textualmente_plausiveis_exigem_revisao():
    b = base([regra("A", descricao_regra="Coca Cola PET"), regra("B", descricao_regra="Coca Cola 2L")])
    r = consulta(b, descricao="Coca Cola PET 2L")
    assert r.regra_priorizada is None
    assert r.nivel_confianca == "REVISAO_NECESSARIA"


def test_prefixo_identificado_e_confianca_reduzida():
    r = consulta(base([regra("P", "1234", "PREFIXO", exige_descricao="NAO")]))
    assert r.regra_priorizada.tipo_correspondencia == "PREFIXO"
    assert r.regra_priorizada.prioridade == (1, 4)
    assert r.nivel_confianca == "MEDIA"


def test_regra_fora_da_vigencia_nao_priorizada():
    r = consultar(base([regra(vigencia_fim=date(2025, 12, 31), exige_descricao="NAO")]), "12345678", data_referencia=HOJE)
    assert r.regra_priorizada is None
    assert r.candidatos_aliquota[0].vigencia == "FORA_DA_VIGENCIA"


@pytest.mark.parametrize(("campo", "valor"), [
    ("exige_descricao", "TALVEZ"), ("aplicacao", "SOMAR"),
    ("escopo_operacao", "EXTERNA"),
])
def test_loader_rejeita_novas_opcoes_invalidas(tmp_path, campo, valor):
    caminho = tmp_path / "opcao_invalida.xlsx"
    linha = beneficio(exige_descricao="SIM", aplicacao="UNICO", grupo_beneficio="")
    linha[campo] = valor
    _planilha_teste(caminho, beneficios=[linha])
    with pytest.raises(ErroBase, match=campo):
        carregar_base(caminho)


def test_planilha_mantem_abas_colunas_e_validacoes_sem_dados_fiscais():
    workbook = openpyxl.load_workbook(CAMINHO_BASE)
    assert {"NCM", "Aliquotas", "Beneficios", "Legislacao"} <= set(workbook.sheetnames)
    assert "exige_descricao" in [c.value for c in workbook["Aliquotas"][1]]
    assert {"exige_descricao", "grupo_beneficio", "aplicacao"} <= {c.value for c in workbook["Beneficios"][1]}
    assert any(v.formula1 == '"SIM,NAO"' for v in workbook["Aliquotas"].data_validations.dataValidation)
    assert any(v.formula1 == '"UNICO,ALTERNATIVO,CUMULATIVO"' for v in workbook["Beneficios"].data_validations.dataValidation)
    for nome in ("NCM", "Aliquotas", "Beneficios", "Legislacao"):
        assert all(cell.value is None for row in workbook[nome].iter_rows(min_row=2) for cell in row)
